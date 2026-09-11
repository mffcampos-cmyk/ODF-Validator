from __future__ import annotations
from pathlib import Path
import openpyxl
from .tables import CodeRow, CodeTable, squash_name as _squash

SKIP_SHEETS = {"DOCUMENT_CONTROL", "CHANGE_LOG", "CONTENTS"}
ID_HEADERS = ("Id", "Code")
SPORT_CODES_SHEET = "SPORT_CODES"


def _norm(h) -> str:
    return str(h).strip().replace(" ", "_")


def _pick_id_index(headers: list[str], sheet_title: str) -> int:
    """Choose the column that holds the code (table key).

    Preference order:
      1. an explicit 'Id'/'Code' header,
      2. a column whose name matches the sheet name (e.g. the RECORD_TYPE
         sheet's 'Recordtype' column -- there is no 'Id'/'Code' there, and
         column 0 is 'Discipline', which is NOT the code),
      3. column 0 as a last resort.
    """
    for i, h in enumerate(headers):
        if h in ID_HEADERS:
            return i
    target = _squash(sheet_title)
    for i, h in enumerate(headers):
        if _squash(h) == target:
            return i
    return 0


def _load_generic_sheet(ws) -> CodeTable | None:
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return None
    if header is None:
        return None
    headers = [_norm(h) if h is not None else f"col{i}" for i, h in enumerate(header)]
    id_idx = _pick_id_index(headers, ws.title)
    table = CodeTable(ws.title)
    for raw in rows:
        if raw is None or id_idx >= len(raw) or raw[id_idx] is None:
            continue
        code = str(raw[id_idx]).strip()
        if not code:
            continue
        fields = {headers[i]: ("" if v is None else str(v))
                  for i, v in enumerate(raw) if i != id_idx and i < len(headers)}
        table.add_row(CodeRow(id=code, fields=fields))
    return table


def _load_sport_codes(ws) -> list[CodeTable]:
    """SPORT_CODES is partitioned by Discipline + Code_Entity (e.g. @TeamType).
    Expand it into one table per (entity, discipline), named SC<Entity>@<Discipline>
    (e.g. 'SC@TeamType@ARC'), keyed by the Code column. This lets rules reference
    SC@ sport codes the way the documentation does."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    idx = {name: headers.index(name) for name in
           ("Discipline", "Code_Entity", "Code") if name in headers}
    if not {"Discipline", "Code_Entity", "Code"} <= idx.keys():
        # unexpected layout -> fall back to a generic table
        t = _load_generic_sheet(ws)
        return [t] if t else []
    eng = headers.index("ENG_Description") if "ENG_Description" in headers else None
    fra = headers.index("FRA_Description") if "FRA_Description" in headers else None
    tables: dict[str, CodeTable] = {}
    for raw in rows[1:]:
        if raw is None:
            continue
        disc = raw[idx["Discipline"]]
        ent = raw[idx["Code_Entity"]]
        code = raw[idx["Code"]]
        if disc is None or ent is None or code is None:
            continue
        disc = str(disc).strip(); ent = str(ent).strip(); code = str(code).strip()
        if not (disc and ent and code):
            continue
        name = f"SC{ent}@{disc}"          # e.g. SC@TeamType@ARC
        fields = {}
        if eng is not None and eng < len(raw) and raw[eng] is not None:
            fields["ENG_Description"] = str(raw[eng])
        if fra is not None and fra < len(raw) and raw[fra] is not None:
            fields["FRA_Description"] = str(raw[fra])
        tables.setdefault(name, CodeTable(name)).add_row(CodeRow(id=code, fields=fields))
    return list(tables.values())


def load_excel_codes(path: Path) -> list[CodeTable]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables: list[CodeTable] = []
    try:
        for ws in wb.worksheets:
            if ws.title in SKIP_SHEETS:
                continue
            if ws.title == SPORT_CODES_SHEET:
                tables.extend(_load_sport_codes(ws))
            else:
                t = _load_generic_sheet(ws)
                if t is not None:
                    tables.append(t)
    finally:
        wb.close()
    return tables
